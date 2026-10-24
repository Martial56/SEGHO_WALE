"""Écrans des états (lot 4) : sommaire, affichage, impression, export Excel."""
import io
from datetime import date, timedelta
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import Http404, HttpResponse
from django.shortcuts import render
from django.utils import timezone

from .etats import ETATS, GROUPES, peut_voir

PERIODES = [('jour', "Aujourd'hui"), ('semaine', 'Cette semaine'), ('mois', 'Ce mois'), ('annee', 'Cette année')]


def bornes(periode, aujourdhui):
    if periode == 'jour':
        return aujourdhui, aujourdhui
    if periode == 'semaine':
        return aujourdhui - timedelta(days=aujourdhui.weekday()), aujourdhui
    if periode == 'annee':
        return aujourdhui.replace(month=1, day=1), aujourdhui
    return aujourdhui.replace(day=1), aujourdhui


def _date(valeur):
    try:
        return date.fromisoformat(valeur)
    except (TypeError, ValueError):
        return None


@login_required(login_url='login')
def etats(request):
    """Sommaire des états accessibles, classés comme dans la proposition."""
    visibles = [(code, e) for code, e in ETATS.items() if peut_voir(request.user, e)]
    if not visibles:
        raise PermissionDenied
    groupes = [(titre, [v for v in visibles if v[1]['groupe'] == cle]) for cle, titre in GROUPES]
    return render(request, 'comptabilite/etats.html', {'groupes': [g for g in groupes if g[1]]})


@login_required(login_url='login')
def etat(request, code):
    definition = ETATS.get(code)
    if definition is None:
        raise Http404
    if not peut_voir(request.user, definition):
        raise PermissionDenied
    aujourdhui = timezone.localdate()
    periode = request.GET.get('periode') or ('' if request.GET.get('debut') else definition['periode'])
    debut, fin = bornes(periode, aujourdhui) if periode else (None, None)
    debut = _date(request.GET.get('debut')) or debut or aujourdhui.replace(day=1)
    fin = _date(request.GET.get('fin')) or fin or aujourdhui
    if fin < debut:
        debut, fin = fin, debut
    params = {cle: request.GET.get(cle, '').strip() for cle in ('caisse', 'journal', 'compte')}
    params['brouillons'] = request.GET.get('brouillons') == '1'

    resultat = definition['fonction'](debut, fin, params)
    if request.GET.get('format') == 'xlsx':
        return _excel(code, resultat, debut, fin)
    # Chaque cellule avec le genre de sa colonne : le gabarit aligne et
    # formate les montants sans avoir à croiser deux listes.
    for t in resultat['tableaux']:
        genres = [g for _, g in t['colonnes']]
        for l in t['lignes']:
            l['paires'] = list(zip(l['cellules'], genres))
            l['vide_apres'] = len(genres) - len(l['cellules'])
    return render(request, 'comptabilite/etat.html', {
        'code': code, 'definition': definition, 'etat': resultat, 'debut': debut, 'fin': fin,
        'periode': periode, 'periodes': PERIODES, 'params': params,
        'comptable': definition['perm'] != 'caisse', 'imprime_le': timezone.now(),
    })


def _valeur(v):
    if v is None:
        return ''
    if isinstance(v, Decimal):
        return int(v) if v == v.to_integral_value() else float(v)
    return v


def _excel(code, resultat, debut, fin):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from rapports.export import populate_sheet

    wb = Workbook()
    tableaux = resultat['tableaux'] or [{'titre': resultat['titre'], 'colonnes': [('', 'texte')], 'lignes': []}]
    for i, t in enumerate(tableaux):
        ws = wb.active if i == 0 else wb.create_sheet()
        lignes = [[_valeur(c) for c in l['cellules']] for l in t['lignes']]
        titre = f"{resultat['titre']} — {t['titre']}"
        # Un nom d'onglet Excel refuse / \ : * ? [ ] : on le nettoie, et le
        # bandeau reprend le vrai titre (dates comprises).
        onglet = ''.join('-' if ch in '/\\:*?[]' else ch for ch in t['titre'])[:31]
        entete = populate_sheet(ws, onglet, [c[0] for c in t['colonnes']], lignes,
                                debut, fin, timezone.localtime())
        ws.cell(row=1, column=1, value=titre.upper())
        gras = Font(name='Times New Roman', bold=True)
        for j, l in enumerate(t['lignes'], start=entete + 1):
            if l['type'] != 'ligne':
                for cell in ws[j]:
                    cell.font = gras
            for k, (_, genre) in enumerate(t['colonnes'], start=1):
                if genre == 'montant':
                    ws.cell(row=j, column=k).number_format = '#,##0'
    tampon = io.BytesIO()
    wb.save(tampon)
    reponse = HttpResponse(tampon.getvalue(),
                           content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    reponse['Content-Disposition'] = f'attachment; filename="{code}_{debut:%Y%m%d}_{fin:%Y%m%d}.xlsx"'
    return reponse
