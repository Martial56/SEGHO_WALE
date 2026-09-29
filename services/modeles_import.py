"""Classeurs vierges à télécharger avant d'importer des prestations ou des
catégories.

Un import échoue presque toujours pour la même raison : la personne a inventé
les colonnes, ou tapé « Gynécologie » là où le fichier attend `GYN`. Le modèle
répond aux deux — il porte exactement les colonnes lues par l'import, dans le
même ordre, et propose en liste déroulante les codes qui existent réellement
dans cette base.

Les listes sont construites à la demande, jamais figées : un département créé
ce matin apparaît dans le modèle téléchargé cet après-midi.
"""
import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from django.http import HttpResponse


#: Nombre de lignes couvertes par les listes déroulantes. Au-delà, la colonne
#: reste saisissable — la validation n'est qu'une aide, pas un verrou.
NB_LIGNES = 500

_REMPLISSAGE_ENTETE = PatternFill('solid', fgColor='2E7D32')
_POLICE_ENTETE = Font(bold=True, color='FFFFFF', size=11)
_BORDURE = Border(
    left=Side(style='thin', color='D0D0D0'), right=Side(style='thin', color='D0D0D0'),
    top=Side(style='thin', color='D0D0D0'), bottom=Side(style='thin', color='D0D0D0'),
)


def _poser_les_colonnes(ws, entetes, exemple):
    """L'en-tête, puis une ligne d'exemple en gris italique.

    L'exemple se supprime d'un coup et ne risque pas d'être pris pour une
    donnée : il ne ressemble pas à ce qu'on saisit.
    """
    for i, titre in enumerate(entetes, 1):
        cellule = ws.cell(row=1, column=i, value=titre)
        cellule.font = _POLICE_ENTETE
        cellule.fill = _REMPLISSAGE_ENTETE
        cellule.alignment = Alignment(horizontal='center', vertical='center')
        cellule.border = _BORDURE
    ws.row_dimensions[1].height = 20

    for i, valeur in enumerate(exemple, 1):
        cellule = ws.cell(row=2, column=i, value=valeur)
        cellule.border = _BORDURE
        cellule.font = Font(italic=True, color='888888')

    for i in range(1, len(entetes) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22


def _poser_les_listes(wb, ws, entetes, listes):
    """Ajoute les listes déroulantes décrites par {colonne: [valeurs]}.

    Les valeurs vivent dans une feuille cachée : Excel n'accepte pas une liste
    littérale de plus de 255 caractères, et un catalogue de catégories la
    dépasse vite.
    """
    feuille = wb.create_sheet('Listes')
    feuille.sheet_state = 'hidden'

    for rang, (colonne, valeurs) in enumerate(listes.items()):
        if colonne not in entetes:
            continue
        valeurs = [v for v in valeurs if v] or ['']
        lettre_source = get_column_letter(rang + 1)
        for i, valeur in enumerate(valeurs, 1):
            feuille[f'{lettre_source}{i}'] = valeur

        validation = DataValidation(
            type='list',
            formula1=f"'Listes'!${lettre_source}$1:${lettre_source}${len(valeurs)}",
            # `allow_blank` et pas d'erreur bloquante : une colonne facultative
            # doit pouvoir rester vide, et un code absent de la liste reste
            # acceptable — l'import le signalera lui-même.
            allow_blank=True, showErrorMessage=False,
        )
        ws.add_data_validation(validation)
        lettre_cible = get_column_letter(entetes.index(colonne) + 1)
        validation.add(f'{lettre_cible}2:{lettre_cible}{NB_LIGNES}')


def classeur_modele(titre, entetes, exemple, listes, nom_fichier):
    """Le classeur prêt à télécharger."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titre

    _poser_les_colonnes(ws, entetes, exemple)
    _poser_les_listes(wb, ws, entetes, listes)

    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    reponse = HttpResponse(
        tampon.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    reponse['Content-Disposition'] = f'attachment; filename="{nom_fichier}"'
    return reponse
