"""Télécharger ce que la liste montre — filtres et regroupement compris.

Le menu « Exporter » sortait la table entière. On filtrait « femmes de plus de
60 ans », on téléchargeait, et on recevait les 47 813 patients : le fichier ne
répondait jamais à la question qu'on venait de poser à l'écran.

Ce module part du jeu **déjà filtré par la vue** — le même queryset que celui
qui alimente la page — et le rend dans le format demandé. Quand un regroupement
est posé, le fichier en porte la structure : un titre par groupe, suivi de ses
lignes, comme à l'écran. Un export à plat redeviendrait muet sur le découpage
qu'on vient justement de choisir.

Les chemins de groupe sont calculés par `core.listing._chemins`, celui-là même
qui range les lignes sous leurs en-têtes dans la page : les deux ne peuvent donc
pas diverger.
"""

import csv
import io
import json
from datetime import datetime

from django.http import HttpResponse
from django.utils import timezone

from core.listing import _attribut, _chemins

#: Formats proposés par le menu de chaque liste.
FORMATS = ('json', 'csv', 'xlsx')

#: Taille des paquets lus en base. Un export de plusieurs dizaines de milliers
#: de lignes ne doit pas charger tout le jeu d'un coup.
LOT = 2000

_ENTETE_FOND = '1F6E8C'
_GROUPE_FOND = 'E8F1F6'


class Colonne:
    """Une colonne d'export : un en-tête, et de quoi en tirer la valeur.

    `valeur` est soit un chemin de champ (`'patient__nom'`, traversant les
    relations et les JSONField comme ailleurs dans le moteur de listes), soit un
    appelable qui reçoit l'objet.
    """

    def __init__(self, entete, valeur, largeur=None):
        self.entete = entete
        self.largeur = largeur
        if callable(valeur):
            self._lire = valeur
        else:
            self._lire = lambda objet, chemin=valeur: _attribut(objet, chemin)

    def lire(self, objet):
        valeur = self._lire(objet)
        return '' if valeur is None else valeur


def _valeur_cellule(valeur):
    """Valeur écrite dans une cellule Excel.

    Un booléen y sortirait « VRAI »/« FAUX » selon la langue du tableur, ou
    « TRUE » si Excel n'est pas francisé : on écrit « Oui »/« Non », comme
    partout ailleurs dans l'application et comme dans les deux autres formats.
    """
    if valeur is None:
        return ''
    if isinstance(valeur, bool):
        return 'Oui' if valeur else 'Non'
    return _heure_locale(valeur)


def _heure_locale(valeur):
    """Horodatage ramené au fuseau du centre, sans fuseau attaché.

    Les dates sont stockées en UTC. Rendues telles quelles, un soin de 7 h du
    matin sortirait à une autre heure que celle lue à l'écran — et Excel, qui
    ne sait pas représenter un fuseau, refuse purement et simplement
    d'enregistrer le classeur.
    """
    if isinstance(valeur, datetime) and timezone.is_aware(valeur):
        return timezone.localtime(valeur).replace(tzinfo=None)
    return valeur


def _texte(valeur):
    """Rendu d'une valeur pour les formats sans typage (CSV, JSON)."""
    if valeur is None or valeur == '':
        return ''
    valeur = _heure_locale(valeur)
    if isinstance(valeur, datetime):
        return valeur.strftime('%d/%m/%Y %H:%M')
    if hasattr(valeur, 'strftime'):
        return valeur.strftime('%d/%m/%Y')
    if isinstance(valeur, bool):
        return 'Oui' if valeur else 'Non'
    return str(valeur)


def _titre_groupe(libelle, nombre):
    return f'{libelle} ({nombre})'


def _totaux_par_prefixe(paquets):
    """Nombre de lignes sous chaque niveau de groupe, feuilles comprises.

    Un groupe parent ne porte aucune ligne en propre : son compte est la somme
    de ceux de ses enfants. Sans ce relevé, l'en-tête « Dr Assi Serge » ne
    pourrait annoncer que le contenu de son dernier sous-groupe.
    """
    totaux = {}
    for chemin, lignes in paquets:
        for rang in range(1, len(chemin) + 1):
            prefixe = chemin[:rang]
            totaux[prefixe] = totaux.get(prefixe, 0) + len(lignes)
    return totaux


def parcours(paquets):
    """Suite des lignes à écrire : les en-têtes de groupe, puis leurs lignes.

    Les deux niveaux d'un regroupement « médecin puis type de consultation »
    sortaient écrasés sur un seul titre, « ASSI Serge › CONSULTATION ADULTE »,
    réécrit en entier à chaque sous-groupe. Ils occupent désormais une ligne
    chacun, emboîtés : le médecin une fois, ses types en dessous, comme à
    l'écran.

    Rend `('groupe', niveau, libellé, nombre)` ou `('ligne', valeurs)`.
    """
    totaux = _totaux_par_prefixe(paquets)
    precedent = ()
    for chemin, lignes in paquets:
        # Les paquets sont triés sur le libellé : les frères se suivent, et le
        # préfixe commun avec le chemin précédent est déjà écrit.
        commun = 0
        while (commun < len(precedent) and commun < len(chemin)
               and precedent[commun] == chemin[commun]):
            commun += 1
        for niveau in range(commun, len(chemin)):
            yield ('groupe', niveau, _texte(chemin[niveau]) or '—',
                   totaux[chemin[:niveau + 1]])
        for ligne in lignes:
            yield ('ligne', ligne)
        precedent = chemin


def blocs(qs, colonnes, dims=()):
    """Lignes rendues, rangées par groupe quand un regroupement est posé.

    Rend une liste de `(chemin, lignes)`. Sans regroupement, un seul bloc de
    chemin vide. Une ligne dont la dimension rend plusieurs libellés apparaît
    dans plusieurs groupes, comme à l'écran.
    """
    if not dims:
        return [((), [[c.lire(o) for c in colonnes] for o in qs.iterator(chunk_size=LOT)])]

    groupes = {}
    for objet in qs.iterator(chunk_size=LOT):
        ligne = [c.lire(objet) for c in colonnes]
        for chemin in _chemins(objet, dims):
            groupes.setdefault(chemin, []).append(ligne)
    # Tri sur le libellé : l'ordre des en-têtes reste lisible (alphabétique,
    # chronologique sur une date formatée en tête) et surtout déterministe.
    return sorted(groupes.items(), key=lambda kv: tuple(_texte(v) for v in kv[0]))


# ── Rendus ──────────────────────────────────────────────────────────────────

def _reponse_csv(nom, colonnes, paquets, separateur=';'):
    """CSV de la sélection.

    Le séparateur par défaut est le point-virgule, celui qu'attend un Excel
    francophone. Les listes dont le CSV se réimporte dans l'application passent
    la virgule : leur lecteur est un `csv.DictReader` sans séparateur déclaré,
    et un fichier en point-virgule y arriverait en une seule colonne.
    """
    reponse = HttpResponse(content_type='text/csv; charset=utf-8')
    reponse['Content-Disposition'] = f'attachment; filename="{nom}.csv"'
    reponse.write('﻿')                       # BOM : accents lisibles dans Excel
    graveur = csv.writer(reponse, delimiter=separateur)
    graveur.writerow([c.entete for c in colonnes])
    for evenement in parcours(paquets):
        if evenement[0] == 'groupe':
            _, niveau, libelle, nombre = evenement
            # Deux espaces par niveau : un fichier texte n'a pas de retrait,
            # et sans eux rien ne distingue un sous-groupe de son parent.
            graveur.writerow(['  ' * niveau + _titre_groupe(libelle, nombre)])
        else:
            graveur.writerow([_cellule_sure(_texte(v)) for v in evenement[1]])
    return reponse


#: Caractères par lesquels Excel reconnaît une formule. Une valeur qui commence
#: par l'un d'eux part précédée d'une apostrophe, sinon un nom de patient
#: commençant par « = » serait exécuté à l'ouverture du fichier.
_AMORCES_FORMULE = ('=', '+', '-', '@', '\t', '\r')


def _cellule_sure(texte):
    return "'" + texte if texte[:1] in _AMORCES_FORMULE else texte


def _noeuds(paquets, entetes):
    """Groupes emboîtés, chacun avec le nombre de lignes qu'il porte.

    Un tableau plat de feuilles ne saurait pas dire combien de rendez-vous
    tient un médecin : l'information n'existe qu'au-dessus des feuilles.
    """
    racine, index = [], {}
    for chemin, lignes in paquets:
        freres = racine
        for rang, valeur in enumerate(chemin):
            prefixe = chemin[:rang + 1]
            noeud = index.get(prefixe)
            if noeud is None:
                noeud = {'groupe': _texte(valeur) or '—', 'nombre': 0}
                index[prefixe] = noeud
                freres.append(noeud)
            noeud['nombre'] += len(lignes)
            if rang < len(chemin) - 1:
                freres = noeud.setdefault('sous_groupes', [])
            else:
                noeud.setdefault('lignes', []).extend(
                    dict(zip(entetes, (_texte(v) for v in ligne))) for ligne in lignes)
    return racine


def _reponse_json(nom, colonnes, paquets, groupe):
    entetes = [c.entete for c in colonnes]
    if groupe:
        data = _noeuds(paquets, entetes)
    else:
        data = [dict(zip(entetes, (_texte(v) for v in ligne)))
                for _chemin, lignes in paquets for ligne in lignes]
    reponse = HttpResponse(
        json.dumps(data, ensure_ascii=False, indent=2, default=str),
        content_type='application/json')
    reponse['Content-Disposition'] = f'attachment; filename="{nom}.json"'
    return reponse


def _reponse_xlsx(nom, titre_feuille, colonnes, paquets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = (titre_feuille or nom)[:31]

    feuille.append([c.entete for c in colonnes])
    fond = PatternFill(start_color=_ENTETE_FOND, end_color=_ENTETE_FOND, fill_type='solid')
    for cellule in feuille[1]:
        cellule.fill, cellule.font = fond, Font(color='FFFFFF', bold=True)
        cellule.alignment = Alignment(horizontal='center')

    fond_groupe = PatternFill(start_color=_GROUPE_FOND, end_color=_GROUPE_FOND, fill_type='solid')
    # Largeurs calculées au fil de l'écriture, sur les seules lignes de
    # données : relire la feuille ensuite ferait entrer les titres de groupe —
    # longs, et seuls dans la première colonne — dans le calcul, qui étirerait
    # cette colonne pour rien.
    largeurs = [len(c.entete) for c in colonnes]

    # Le rang est tenu ici, et non lu dans `feuille.max_row` : cette propriété
    # parcourt toutes les cellules du classeur à chaque appel. Appelée une fois
    # par ligne datée — donc 12 000 fois sur un export de rendez-vous, chacune
    # balayant 180 000 cellules — elle coûtait à elle seule près de trois
    # minutes, pour un fichier que le CSV rend en neuf secondes.
    rang = 1                                 # la ligne d'en-tête est écrite

    for evenement in parcours(paquets):
        if evenement[0] == 'groupe':
            _, niveau, libelle, nombre = evenement
            feuille.append([_titre_groupe(libelle, nombre)])
            rang += 1
            cellule = feuille.cell(row=rang, column=1)
            # Le retrait dit l'emboîtement, comme à l'écran. Le premier niveau
            # garde sa bande de fond : sur une feuille de cinquante colonnes,
            # c'est elle qui donne les grandes coupures du regard.
            cellule.alignment = Alignment(indent=niveau)
            cellule.font = Font(bold=True, color='14536B')
            if niveau == 0:
                # La bande court sur toute la largeur : sans fond sur les
                # cellules vides, la couleur s'arrêterait au texte.
                for colonne in range(1, len(colonnes) + 1):
                    feuille.cell(row=rang, column=colonne).fill = fond_groupe
            continue

        ligne = evenement[1]
        feuille.append([_valeur_cellule(v) for v in ligne])
        rang += 1
        for colonne, valeur in enumerate(ligne, start=1):
            largeurs[colonne - 1] = max(largeurs[colonne - 1], len(_texte(valeur)))
            # Une date écrite sans format sort en ISO ; on la montre comme
            # partout ailleurs dans l'application.
            if not hasattr(valeur, 'strftime'):
                continue
            feuille.cell(row=rang, column=colonne).number_format = (
                'DD/MM/YYYY HH:MM' if isinstance(valeur, datetime) else 'DD/MM/YYYY')

    for index, colonne in enumerate(colonnes, start=1):
        lettre = feuille.cell(row=1, column=index).column_letter
        feuille.column_dimensions[lettre].width = (
            colonne.largeur or min(max(largeurs[index - 1] + 3, 10), 50))
    feuille.freeze_panes = 'A2'

    tampon = io.BytesIO()
    classeur.save(tampon)
    reponse = HttpResponse(
        tampon.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    reponse['Content-Disposition'] = f'attachment; filename="{nom}.xlsx"'
    return reponse


def contexte(request, selection, url, libelle, nombre=None):
    """Variables attendues par `includes/listing/menu_export.html`.

    `nombre` évite un `COUNT` que la vue a souvent déjà fait : avec un
    regroupement le paginateur compte les groupes et non les lignes, la vue
    passe alors None et on compte ici.
    """
    from core.listing import parametres_export
    return {
        'export_url': url,
        'export_libelle': libelle,
        'export_qs': parametres_export(request),
        'nb_selection': selection.qs.count() if nombre is None else nombre,
    }


def repondre(fmt, nom, colonnes, qs, dims=(), titre_feuille=None, separateur=';'):
    """Réponse HTTP portant le jeu filtré, dans le format demandé.

    `fmt` inconnu retombe sur Excel plutôt que de refuser : le lien vient d'un
    menu, pas d'une saisie, et un format mal orthographié n'est pas une raison
    de ne rien rendre.
    """
    paquets = blocs(qs, colonnes, dims)
    if fmt == 'csv':
        return _reponse_csv(nom, colonnes, paquets, separateur)
    if fmt == 'json':
        return _reponse_json(nom, colonnes, paquets, bool(dims))
    return _reponse_xlsx(nom, titre_feuille, colonnes, paquets)
