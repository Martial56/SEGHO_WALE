#!/usr/bin/env python3
"""Teste l'algorithme de détection des doublons sur un fichier Excel de patients.

Utilise EXACTEMENT les fonctions de patients/doublons.py (normalisation,
phonétique, similarité des noms, des dates, des téléphones, poids, seuils et
calcul du score) : ce que vous voyez dans le rapport est ce que l'application
fera. Ne touche ni à la base de données ni au fichier d'origine.

Le rapport Excel produit contient :
    Résumé ........ les chiffres clés et le mode d'emploi (une page)
    À fusionner ... les groupes de fiches quasi certainement identiques
                    (une fiche « à garder » est suggérée dans chaque groupe)
    À vérifier .... les paires douteuses, fiche A et fiche B côte à côte
    Patients analysés, et « Test d'efficacité » si --simuler est utilisé.

Dates « par défaut » : quand un fichier ne contient que des âges, la date de
naissance est calculée et beaucoup de fiches tombent sur le même jour/mois
(ex. 30/09 pour 89 % d'une base). Le script détecte ces jours-là tout seul et
les traite comme « année seulement » : ils ne comptent plus comme preuve.

Utilisation (le script doit être dans le dossier du projet SEGHO_WALE) :

    Double-cliquez sur le fichier, ou :  python tester_doublons_excel.py
    → une fenêtre s'ouvre : « Parcourir… », choisir l'Excel, « Lancer l'analyse ».

En ligne de commande, si vous préférez :

    python tester_doublons_excel.py patients.xlsx
    python tester_doublons_excel.py patients.xlsx --simuler 200
    python tester_doublons_excel.py patients.xlsx --sortie rapport.xlsx --feuille Patients

Colonnes reconnues automatiquement (l'ordre et la casse n'ont pas d'importance) :
    nom (seul, ou « NOM Prénoms » dans une seule colonne), prenoms,
    date_naissance OU age (34, ou « 34Ans2Mois5Jours » comme l'export de l'app),
    sexe / genre, telephone / mobile, telephone2, code / code_identifiant.
Si un nom de colonne n'est pas reconnu : --col-nom "Nom du patient", etc.

--simuler N : fabrique N faux doublons à partir de vraies fiches (fautes,
variantes phonétiques, prénoms inversés, autre téléphone, jour/mois inversés…)
et mesure combien l'algorithme en retrouve. Le rapport principal, lui, ne
porte que sur vos vraies données.
"""

import argparse
import os
import random
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta

# ── Charger patients/doublons.py sans démarrer tout Django ──────────────────
ICI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ICI)
try:
    from django.conf import settings  # noqa: E402
    if not settings.configured:
        settings.configure()
except ImportError:
    # Django absent de ce Python (double-clic hors de l'environnement du
    # projet) : doublons.py n'en utilise que deux noms, on les simule.
    import types
    _conf = types.ModuleType('django.conf')
    _conf.settings = types.SimpleNamespace(configured=True)
    _models = types.ModuleType('django.db.models')
    _models.Q = object
    for _nom, _mod in {'django': types.ModuleType('django'), 'django.conf': _conf,
                       'django.db': types.ModuleType('django.db'), 'django.db.models': _models}.items():
        sys.modules[_nom] = _mod
if not os.path.isfile(os.path.join(ICI, 'patients', 'doublons.py')):
    _msg = ("Ce script doit être placé à la racine du projet SEGHO_WALE "
            "(à côté de manage.py), avec le fichier patients/doublons.py.")
    try:
        import tkinter.messagebox
        tkinter.Tk().withdraw()
        tkinter.messagebox.showerror('Fichier manquant', _msg)
    except Exception:
        print(_msg)
    sys.exit(1)
from patients import doublons as D  # noqa: E402
from functools import lru_cache  # noqa: E402

# Mêmes fonctions, mises en cache : sur un gros fichier les mêmes noms
# reviennent des milliers de fois (KOUASSI, AMA…), inutile de tout recalculer.
D.cle_phonetique = lru_cache(maxsize=200_000)(D.cle_phonetique)
D.similarite_mot = lru_cache(maxsize=2_000_000)(D.similarite_mot)

try:
    from openpyxl import Workbook, load_workbook  # noqa: E402
except ImportError:
    _msg = "Le module openpyxl manque. Installez-le avec :\n\n    pip install openpyxl"
    try:
        import tkinter.messagebox
        tkinter.Tk().withdraw()
        tkinter.messagebox.showerror('Module manquant', _msg)
    except Exception:
        print(_msg)
    sys.exit(1)
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

# ── Lecture du fichier ──────────────────────────────────────────────────────

SYNONYMES = {
    'code': ['code_identifiant', 'code_patient', 'code', 'identifiant', 'id', 'ndossier', 'numero',
             'numerodossier', 'matricule'],
    'nom': ['nom', 'noms', 'nomcomplet', 'nometprenoms', 'nomprenoms', 'patient', 'nompatient'],
    'prenoms': ['prenoms', 'prenom'],
    'date_naissance': ['date_naissance', 'datenaissance', 'datedenaissance', 'naissance', 'ddn', 'nele',
                       'datenais', 'dob'],
    'age': ['age', 'ageans'],
    'sexe': ['sexe', 'genre', 'sex'],
    'telephone': ['telephone', 'mobile', 'tel', 'contact', 'portable', 'telephone1', 'mobile1', 'cellulaire'],
    'telephone2': ['telephone2', 'mobile2', 'tel2', 'contact2', 'autretelephone'],
}


def _cle_entete(s):
    s = unicodedata.normalize('NFKD', str(s or ''))
    s = ''.join(c for c in s if not unicodedata.combining(c)).lower()
    s = s.replace('°', '')
    return re.sub(r'[^a-z0-9_]', '', s.replace(' ', ''))


def detecter_colonnes(entetes, forcees):
    cles = [_cle_entete(e) for e in entetes]
    trouve = {}
    for champ, noms in SYNONYMES.items():
        if forcees.get(champ):
            cible = _cle_entete(forcees[champ])
            if cible not in cles:
                raise ValueError(f'Colonne « {forcees[champ]} » introuvable. Colonnes du fichier : {entetes}')
            trouve[champ] = cles.index(cible)
            continue
        for n in noms:
            if n in cles and cles.index(n) not in trouve.values():
                trouve[champ] = cles.index(n)
                break
    return trouve


def lire_age(v, ref):
    """(date de naissance estimée, précise ?) depuis « 34 » ou « 34Ans2Mois5Jours »."""
    s = str(v or '').strip()
    if not s:
        return None, False
    if re.fullmatch(r'\d+([.,]\d+)?', s):
        ans = int(float(s.replace(',', '.')))
        return _reculer(ref, ans * 12, 0), False
    a = re.search(r'(\d+)\s*Ans?', s, re.I)
    m = re.search(r'(\d+)\s*Mois', s, re.I)
    j = re.search(r'(\d+)\s*Jours?', s, re.I)
    if not (a or m or j):
        return None, False
    mois = (int(a.group(1)) if a else 0) * 12 + (int(m.group(1)) if m else 0)
    return _reculer(ref, mois, int(j.group(1)) if j else 0), bool(j)


def _reculer(ref, mois, jours):
    annee, m = ref.year, ref.month - mois
    while m <= 0:
        m += 12
        annee -= 1
    jour = ref.day
    while True:
        try:
            base = date(annee, m, jour)
            break
        except ValueError:
            jour -= 1
    return base - timedelta(days=jours)


def lire_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or '').strip()
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d.%m.%Y', '%d/%m/%y', '%Y/%m/%d'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def lire_sexe(v):
    s = _cle_entete(v)
    return {'m': 'M', 'masculin': 'M', 'homme': 'M', 'h': 'M', 'male': 'M',
            'f': 'F', 'feminin': 'F', 'femme': 'F', 'female': 'F'}.get(s, '')


def lire_fichier(chemin, feuille, forcees, ref):
    wb = load_workbook(chemin, read_only=True, data_only=True)
    ws = wb[feuille] if feuille else wb.worksheets[0]
    lignes = list(ws.iter_rows(values_only=True))
    # La ligne d'en-têtes : la première des 10 premières où l'on reconnaît au moins 2 colonnes.
    idx_entete, cols = None, {}
    for i, ligne in enumerate(lignes[:10]):
        c = detecter_colonnes([x for x in ligne], forcees)
        if len(c) >= 2 and 'nom' in c:
            idx_entete, cols = i, c
            break
    if idx_entete is None:
        raise ValueError("Impossible de trouver la colonne du nom. Précisez-la avec --col-nom \"...\".\n"
                 f"Premières lignes lues : {lignes[:3]}")
    entetes = [str(x or '') for x in lignes[idx_entete]]

    patients, ignores = [], []
    nom_seul = 'prenoms' not in cols
    for n, ligne in enumerate(lignes[idx_entete + 1:], start=idx_entete + 2):
        val = lambda champ: ligne[cols[champ]] if champ in cols and cols[champ] < len(ligne) else None  # noqa
        brut_nom = str(val('nom') or '').strip()
        if not brut_nom:
            if any(x not in (None, '') for x in ligne):
                ignores.append((n, 'nom vide'))
            continue
        if nom_seul:
            parts = brut_nom.split(' ', 1)
            nom, prenoms = parts[0], (parts[1] if len(parts) > 1 else '')
        else:
            nom, prenoms = brut_nom, str(val('prenoms') or '').strip()

        dn, precise = None, False
        if 'date_naissance' in cols:
            dn = lire_date(val('date_naissance'))
            precise = dn is not None
        if dn is None and 'age' in cols:
            dn, precise = lire_age(val('age'), ref)
        patients.append({
            'ligne': n, 'code': str(val('code') or '').strip(),
            'nom': nom, 'prenoms': prenoms,
            'date_naissance': dn, 'date_precise': precise,
            'affichage_naissance': (dn.strftime('%d/%m/%Y') if dn and precise
                                    else (f'~{ref.year - dn.year} ans' if dn else '')),
            'sexe': lire_sexe(val('sexe')),
            'telephone': str(val('telephone') or '').strip(),
            'telephone2': str(val('telephone2') or '').strip(),
        })
    return patients, ignores, entetes, cols


# ── Dates par défaut ────────────────────────────────────────────────────────

PART_DATE_PAR_DEFAUT = 0.02   # un jour/mois porté par plus de 2 % des fiches n'est pas une vraie date


def marquer_dates_par_defaut(patients, part=PART_DATE_PAR_DEFAUT):
    """Un jour/mois qui revient sur des milliers de fiches (30/09, 01/01…) n'est
    pas une date de naissance : c'est un artefact (âge converti en date, saisie
    « au 1er janvier »). Ces dates ne valent que par leur année."""
    precis = [p for p in patients if p['date_naissance'] and p['date_precise']]
    if not precis:
        return []
    cpt = defaultdict(int)
    for p in precis:
        cpt[(p['date_naissance'].day, p['date_naissance'].month)] += 1
    defauts = {k for k, n in cpt.items() if n / len(precis) >= part} | set(D.DATES_PAR_DEFAUT)
    for p in precis:
        d = p['date_naissance']
        if (d.day, d.month) in defauts:
            p['date_precise'] = False
            p['affichage_naissance'] = f'≈ {d.year}'
    return sorted(((f'{j:02d}/{m:02d}', cpt[(j, m)] / len(precis)) for j, m in defauts if cpt[(j, m)]),
                  key=lambda x: -x[1])


# ── Score (calcul central dans patients/doublons.py) ────────────────────────

def score_paire(a, b):
    s_noms = D.similarite_noms(a['nom'], a['prenoms'], b['nom'], b['prenoms'])
    if a['date_precise'] and b['date_precise']:
        s_date = D.similarite_dates(a['date_naissance'], b['date_naissance'])
    else:
        s_date = D.similarite_dates_approx(a['date_naissance'], b['date_naissance'])
    meme_tel = bool(a['_tels'] & b['_tels'])
    score, sexe_diff = D.score_depuis_similarites(s_noms, s_date, meme_tel, a['sexe'], b['sexe'])
    return score, s_noms, s_date, meme_tel, sexe_diff


def raisons(s_noms, s_date, meme_tel, sexe_diff):
    r = []
    if s_noms >= 0.97:
        r.append('même nom')
    elif s_noms >= 0.85:
        r.append('nom très proche')
    else:
        r.append('nom assez proche')
    if s_date == 1.0:
        r.append('même date de naissance')
    elif s_date >= 0.6:
        r.append('même année de naissance')
    elif s_date > 0:
        r.append('âge à un an près')
    elif s_date == 0:
        r.append('date différente')
    if meme_tel:
        r.append('même téléphone')
    if sexe_diff:
        r.append('sexe différent')
    return ', '.join(r)


# ── Recherche des paires ────────────────────────────────────────────────────

def preparer(patients):
    for i, p in enumerate(patients):
        p['_i'] = i
        p['_tels'] = {t for t in (D.normaliser_telephone(p['telephone']),
                                  D.normaliser_telephone(p['telephone2'])) if t}
        p['_mots'] = sorted({m for m in (D.cle_phonetique(x) for x in D.mots(p['nom']) + D.mots(p['prenoms']))
                             if len(m) >= 2})


def _cles(p):
    """Clés de « blocage » : deux fiches ne sont comparées que si elles
    partagent au moins une clé. Chaque clé combine deux indices, pour que les
    noms très courants (KOUASSI, KONE…) ne forment pas des paquets géants :
      - le même téléphone ;
      - deux mots du nom en commun (dans n'importe quel ordre) ;
      - un mot du nom + l'année de naissance (à un an près) ;
      - le début d'un mot du nom + la date de naissance exacte (rattrape les fautes).
    """
    cles = set()
    for t in p['_tels']:
        cles.add('T' + t[-8:])
    mots_ = p['_mots']
    for x in range(len(mots_)):
        for y in range(x + 1, len(mots_)):
            cles.add('W' + mots_[x] + '|' + mots_[y])
    d = p['date_naissance']
    if d:
        for m in mots_:
            cles.add(f'Y{m}|{d.year}')
            cles.add(f'Y{m}|{d.year + 1}')
            if p['date_precise']:
                cles.add(f'D{m[:3]}|{d.isoformat()}')
                # jour et mois inversés
                if d.day <= 12:
                    cles.add(f'D{m[:3]}|{d.year}-{d.day:02d}-{d.month:02d}')
    return cles


def paires_scorees(patients, seuil, progression=None):
    """Compare chaque fiche à ses seuls candidats et ne garde que les paires
    au-dessus du seuil. La mémoire reste faible même sur un gros fichier."""
    blocs = defaultdict(list)
    cles_de = []
    for p in patients:
        c = _cles(p)
        cles_de.append(c)
        for k in c:
            blocs[k].append(p['_i'])
    gros = sum(1 for v in blocs.values() if len(v) > 1000)
    paires = []
    n = len(patients)
    for a in range(n):
        candidats = set()
        for k in cles_de[a]:
            for b in blocs[k]:
                if b > a:
                    candidats.add(b)
        pa = patients[a]
        for b in candidats:
            s, sn, sd, tel, sx = score_paire(pa, patients[b])
            if s >= seuil:
                paires.append({'a': a, 'b': b, 'score': s, 'noms': sn, 'date': sd, 'tel': tel,
                               'sexe_diff': sx, 'raisons': raisons(sn, sd, tel, sx)})
        if progression and (a % 500 == 0 or a == n - 1):
            progression(a + 1, n)
    return paires, gros


TAILLE_MAX_GROUPE = 8   # au-delà, ce n'est plus « la même personne saisie plusieurs fois »


def regrouper(patients, paires, seuil):
    """Groupes de fiches réellement identiques, SANS effet de chaîne.

    L'ancienne méthode disait « A~B et B~C donc A, B, C ensemble » : à force de
    maillons, 37 000 fiches se retrouvaient dans un seul groupe. Ici, on part des
    paires les plus sûres et on ne fusionne deux groupes que si CHAQUE fiche de
    l'un ressemble à CHAQUE fiche de l'autre (au niveau « quasi certain »), et
    jamais au-delà de TAILLE_MAX_GROUPE fiches. Les liens refusés ne sont pas
    perdus : ils passent dans la liste « À vérifier »."""
    connus = {(p['a'], p['b']): p['score'] for p in paires}
    groupe_de = {}

    def score(x, y):
        k = (x, y) if x < y else (y, x)
        if k not in connus:
            connus[k] = score_paire(patients[k[0]], patients[k[1]])[0]
        return connus[k]

    for p in paires:                      # déjà triées du plus sûr au moins sûr
        if p['score'] < seuil:
            break
        a, b = p['a'], p['b']
        ga, gb = groupe_de.get(a), groupe_de.get(b)
        if ga is not None and ga is gb:
            continue
        ma = ga if ga is not None else [a]
        mb = gb if gb is not None else [b]
        if len(ma) + len(mb) > TAILLE_MAX_GROUPE:
            continue
        if all(score(x, y) >= seuil for x in ma for y in mb):
            fusion = ma + mb
            for x in fusion:
                groupe_de[x] = fusion
    vus, liste = set(), []
    for g in groupe_de.values():
        if id(g) in vus:
            continue
        vus.add(id(g))
        membres = sorted(g)
        mini = min(score(x, y) for i, x in enumerate(membres) for y in membres[i + 1:])
        liste.append({'membres': membres, 'score': mini})
    liste.sort(key=lambda g: (-g['score'], -len(g['membres']), g['membres'][0]))
    for k, g in enumerate(liste, 1):
        g['num'] = k
        g['niveau'] = 'Quasi certain'
        g['garder'] = fiche_a_garder(patients, g['membres'])
    return liste


def fiche_a_garder(patients, membres):
    """Suggestion (à confirmer) : la fiche la plus complète, et à égalité la plus
    ancienne (code le plus petit). Sert de fiche principale lors de la fusion."""
    def cle(i):
        p = patients[i]
        completude = ((1 if p['telephone'] else 0) + (1 if p['telephone2'] else 0)
                      + (1 if p['date_precise'] else 0) + (1 if p['sexe'] else 0)
                      + min(len(p['prenoms'].split()), 3))
        return (-completude, p['code'] or 'ZZZ', p['ligne'])
    return min(membres, key=cle)


def analyser(patients, seuil, progression=None):
    preparer(patients)
    paires, _ = paires_scorees(patients, seuil, progression)
    paires.sort(key=lambda p: p['score'], reverse=True)
    groupes = regrouper(patients, paires, D.SEUIL_BLOCAGE)
    groupe_de = {i: g['num'] for g in groupes for i in g['membres']}
    # Les paires déjà réunies dans un groupe n'ont pas besoin d'être revues
    # une deuxième fois ; toutes les autres sont à vérifier.
    a_verifier = [p for p in paires
                  if not (p['a'] in groupe_de and groupe_de.get(p['a']) == groupe_de.get(p['b']))]
    return paires, groupes, a_verifier


def niveau(score):
    if score >= D.SEUIL_BLOCAGE:
        return 'Quasi certain'
    if score >= D.SEUIL_AVERTISSEMENT:
        return 'À vérifier'
    return 'Faible'


# ── Simulation : mesurer ce que l'algorithme retrouve ───────────────────────

VOYELLES = 'AEIOU'
SUBSTITUTIONS = [('OUA', 'WA'), ('OU', 'W'), ('Y', 'I'), ('I', 'Y'), ('C', 'K'), ('K', 'C'),
                 ('PH', 'F'), ('SS', 'S'), ('S', 'SS'), ('N', 'NN'), ('NN', 'N'), ('E', 'É'),
                 ('A', 'AH'), ('O', 'AU')]


def _faute(mot, rnd):
    if len(mot) < 4:
        return mot + mot[-1]
    i = rnd.randrange(1, len(mot) - 1)
    choix = rnd.choice(['suppr', 'double', 'inverse', 'remplace'])
    if choix == 'suppr':
        return mot[:i] + mot[i + 1:]
    if choix == 'double':
        return mot[:i] + mot[i] + mot[i:]
    if choix == 'inverse':
        return mot[:i] + mot[i + 1] + mot[i] + mot[i + 2:]
    return mot[:i] + rnd.choice('ABCDEIKLMNOSTU') + mot[i + 1:]


def _phonetique(mot, rnd):
    m = mot.upper()
    possibles = [(a, b) for a, b in SUBSTITUTIONS if a in m]
    if not possibles:
        return _faute(mot, rnd)
    a, b = rnd.choice(possibles)
    return m.replace(a, b, 1)


def _autre_tel(rnd):
    return '07' + ''.join(rnd.choice('0123456789') for _ in range(8))


VARIANTES = {
    'Faute de frappe dans le nom': lambda p, r: {**p, 'nom': _faute(p['nom'], r)},
    'Variante phonétique (Kouassi/Kwassi…)': lambda p, r: {**p, 'nom': _phonetique(p['nom'], r)},
    'Prénoms dans un autre ordre': lambda p, r: {**p, 'prenoms': ' '.join(reversed(p['prenoms'].split()))},
    'Nom et prénoms inversés': lambda p, r: {**p, 'nom': p['prenoms'].split(' ')[0] if p['prenoms'] else p['nom'],
                                             'prenoms': p['nom']},
    'Un prénom oublié': lambda p, r: {**p, 'prenoms': p['prenoms'].split(' ')[0]},
    'Accents / minuscules / apostrophe': lambda p, r: {**p, 'nom': p['nom'].title().replace('E', 'É', 1),
                                                       'prenoms': p['prenoms'].lower()},
    'Autre numéro de téléphone': lambda p, r: {**p, 'telephone': _autre_tel(r), 'telephone2': ''},
    'Jour et mois inversés': lambda p, r: {**p, 'date_naissance': (
        p['date_naissance'].replace(month=p['date_naissance'].day, day=p['date_naissance'].month)
        if p['date_naissance'] and p['date_precise'] and p['date_naissance'].day <= 12
        and p['date_naissance'].day != p['date_naissance'].month
        else p['date_naissance'] + timedelta(days=1) if p['date_naissance'] else None)},
    'Faute + autre téléphone': lambda p, r: {**p, 'nom': _faute(p['nom'], r), 'telephone': _autre_tel(r),
                                             'telephone2': ''},
    'CUMUL : faute + autre tél. + date décalée d\'un mois': lambda p, r: {
        **p, 'nom': _faute(p['nom'], r), 'telephone': _autre_tel(r), 'telephone2': '',
        'date_naissance': p['date_naissance'] + timedelta(days=30) if p['date_naissance'] else None},
    'CUMUL : prénom oublié + autre tél. + phonétique': lambda p, r: {
        **p, 'nom': _phonetique(p['nom'], r), 'prenoms': p['prenoms'].split(' ')[0],
        'telephone': _autre_tel(r), 'telephone2': ''},
    'Phonétique + date approximative au 1er janvier': lambda p, r: {
        **p, 'nom': _phonetique(p['nom'], r),
        'date_naissance': p['date_naissance'].replace(month=1, day=1) if p['date_naissance'] else None},
}


def simuler(patients, n, seed=42):
    rnd = random.Random(seed)
    sources = [p for p in patients if p['prenoms'] and p['date_naissance']]
    if not sources:
        return None
    base = [{k: v for k, v in p.items() if not k.startswith('_')} for p in patients]
    noms = list(VARIANTES)
    attendu = []
    for k in range(n):
        src = rnd.choice(sources)
        type_ = noms[k % len(noms)]
        faux = VARIANTES[type_](dict(base[src['_i']]), rnd)
        faux['code'] = f'SIMUL{k + 1:04d}'
        base.append(faux)
        attendu.append((src['_i'], len(base) - 1, type_))
    # Le faux doublon est-il au moins comparé à l'original (même clé de blocage) ?
    preparer(base)
    resultats = defaultdict(lambda: {'total': 0, 'bloque': 0, 'alerte': 0, 'manque': 0, 'scores': []})
    exemples = []
    for a, b, t in attendu:
        compare = bool(_cles(base[a]) & _cles(base[b]))
        s = score_paire(base[a], base[b])[0]
        r = resultats[t]
        r['total'] += 1
        r['scores'].append(s if compare else 0)
        if not compare:
            s = 0   # jamais comparées : l'algorithme ne pourrait pas la trouver
        if s >= D.SEUIL_BLOCAGE:
            r['bloque'] += 1
        elif s >= D.SEUIL_AVERTISSEMENT:
            r['alerte'] += 1
        else:
            r['manque'] += 1
            if len(exemples) < 40:
                exemples.append((t, base[a], base[b], s))
    return resultats, exemples


# ── Rapport Excel ───────────────────────────────────────────────────────────

POLICE = 'Arial'
F_TITRE = Font(name=POLICE, size=14, bold=True, color='1F4E5F')
F_GRAS = Font(name=POLICE, bold=True)
F_NORMAL = Font(name=POLICE)
F_GRIS = Font(name=POLICE, italic=True, color='666666')
F_ENTETE = Font(name=POLICE, bold=True, color='FFFFFF')
R_ENTETE = PatternFill('solid', start_color='1F6E8C')
R_ROUGE = PatternFill('solid', start_color='F8D7DA')
R_ORANGE = PatternFill('solid', start_color='FFF3CD')
R_SAISIE = PatternFill('solid', start_color='FFFF00')
R_ALT = [PatternFill('solid', start_color='FFFFFF'), PatternFill('solid', start_color='EEF4F7')]
BORD = Border(top=Side(style='thin', color='1F6E8C'))


def _entetes(ws, ligne, titres, largeurs=None):
    for c, t in enumerate(titres, 1):
        cell = ws.cell(ligne, c, t)
        cell.font, cell.fill = F_ENTETE, R_ENTETE
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    if largeurs:
        for c, w in enumerate(largeurs, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(ligne + 1, 1)


def _police(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and (c.font is None or c.font.name != POLICE):
                c.font = Font(name=POLICE, bold=c.font.bold if c.font else False,
                              italic=c.font.italic if c.font else False,
                              color=c.font.color if c.font else None,
                              size=c.font.size if c.font else 11)


def ecrire_rapport(chemin, source, patients, ignores, cols, entetes, paires, groupes, a_verifier,
                   simulation, dates_defaut=(), max_verif=20000):
    wb = Workbook()
    rs = wb.active
    rs.title = 'Résumé'
    wf = wb.create_sheet('À fusionner')
    wv = wb.create_sheet('À vérifier')
    wd = wb.create_sheet('Patients analysés')
    groupe_de = {i: g['num'] for g in groupes for i in g['membres']}

    # ---------- À fusionner : un bloc par groupe, une seule décision par groupe ----------
    _entetes(wf, 1, ['Groupe', 'Que faire de cette fiche ?', 'Même personne ?\n(Oui / Non)', 'Score\nle plus bas',
                     'Ligne Excel', 'Code', 'Nom', 'Prénoms', 'Naissance', 'Sexe', 'Téléphone', 'Téléphone 2',
                     'Commentaire'],
             [8, 22, 16, 10, 9, 16, 20, 28, 12, 6, 16, 16, 30])
    wf.row_dimensions[1].height = 34
    dv = DataValidation(type='list', formula1='"Oui,Non,En partie"', allow_blank=True)
    wf.add_data_validation(dv)
    ligne = 2
    for g in groupes:
        ordre = [g['garder']] + [i for i in g['membres'] if i != g['garder']]
        for k, i in enumerate(ordre):
            p = patients[i]
            valeurs = [g['num'], '★ GARDER cette fiche' if k == 0 else '→ à fusionner dans la fiche ★',
                       None, g['score'] / 100 if k == 0 else None, p['ligne'], p['code'], p['nom'], p['prenoms'],
                       p['affichage_naissance'], p['sexe'], p['telephone'], p['telephone2'], None]
            for c, v in enumerate(valeurs, 1):
                cell = wf.cell(ligne, c, v)
                cell.font = F_NORMAL
                cell.fill = R_ALT[g['num'] % 2]
                if k == 0:
                    cell.border = BORD
            wf.cell(ligne, 4).number_format = '0%'
            wf.cell(ligne, 2).font = Font(name=POLICE, bold=(k == 0), color='1B5E20' if k == 0 else '666666')
            if k == 0:
                wf.cell(ligne, 3).fill = R_SAISIE
                dv.add(wf.cell(ligne, 3))
            ligne += 1
    fin_f = max(ligne - 1, 2)
    if not groupes:
        wf.cell(2, 1, 'Aucun groupe quasi certain détecté.').font = F_GRIS
    wf.auto_filter.ref = f'A1:M{fin_f}'

    # ---------- À vérifier : une paire par ligne, A et B côte à côte ----------
    _entetes(wv, 1, ['N°', 'Score', 'Même personne ?\n(Oui / Non)', 'Pourquoi', 'Ligne A', 'Patient A',
                     'Naissance A', 'Tél. A', 'Ligne B', 'Patient B', 'Naissance B', 'Tél. B'],
             [6, 8, 16, 46, 8, 34, 11, 15, 8, 34, 11, 15])
    wv.row_dimensions[1].height = 34
    dv2 = DataValidation(type='list', formula1='"Oui,Non"', allow_blank=True)
    wv.add_data_validation(dv2)
    affichees = a_verifier[:max_verif]
    for n, p in enumerate(affichees, 2):
        a, b = patients[p['a']], patients[p['b']]
        vals = [n - 1, p['score'] / 100, None, p['raisons'], a['ligne'], f"{a['nom']} {a['prenoms']}",
                a['affichage_naissance'], a['telephone'], b['ligne'], f"{b['nom']} {b['prenoms']}",
                b['affichage_naissance'], b['telephone']]
        for c, v in enumerate(vals, 1):
            wv.cell(n, c, v).font = F_NORMAL
        wv.cell(n, 2).number_format = '0%'
        wv.cell(n, 2).fill = R_ROUGE if p['score'] >= D.SEUIL_BLOCAGE else R_ORANGE
        wv.cell(n, 3).fill = R_SAISIE
        dv2.add(wv.cell(n, 3))
    fin_v = max(len(affichees) + 1, 2)
    if not affichees:
        wv.cell(2, 1, 'Aucune paire douteuse.').font = F_GRIS
    wv.auto_filter.ref = f'A1:L{fin_v}'

    # ---------- Patients analysés ----------
    _entetes(wd, 1, ['Ligne Excel', 'Groupe à fusionner', 'Code', 'Nom', 'Prénoms', 'Naissance (lue)',
                     'Sexe', 'Téléphone', 'Téléphone 2'], [9, 12, 16, 20, 28, 15, 6, 16, 16])
    for n, p in enumerate(patients, 2):
        vals = [p['ligne'], groupe_de.get(p['_i']), p['code'], p['nom'], p['prenoms'],
                p['affichage_naissance'], p['sexe'], p['telephone'], p['telephone2']]
        for c, v in enumerate(vals, 1):
            wd.cell(n, c, v).font = F_NORMAL
    wd.auto_filter.ref = f'A1:I{len(patients) + 1}'

    # ---------- Résumé (une page) ----------
    rs.column_dimensions['A'].width = 58
    for col in 'BCDEFG':
        rs.column_dimensions[col].width = 15
    rs['A1'] = 'Analyse des doublons patients'
    rs['A1'].font = F_TITRE
    rs['A2'] = (f'Fichier : {os.path.basename(source)} — analysé le {datetime.now():%d/%m/%Y à %H:%M}. '
                f"Algorithme : patients/doublons.py (le même que dans l'application).")
    rs['A2'].font = F_GRIS

    en_trop = sum(len(g['membres']) - 1 for g in groupes)
    r = 4
    rs.cell(r, 1, 'Où en est votre base').font = F_GRAS
    chiffres = [
        ('Patients lus dans le fichier', len(patients), '#,##0'),
        ('Groupes de doublons quasi certains (feuille « À fusionner »)', len(groupes), '#,##0'),
        ('Fiches en trop = fiches à supprimer après fusion', en_trop, '#,##0'),
        ('Part de la base concernée', en_trop / max(len(patients), 1), '0.0%'),
        ('Paires douteuses à examiner (feuille « À vérifier »)', len(a_verifier), '#,##0'),
    ]
    for k, (lib, val, fmt) in enumerate(chiffres, r + 1):
        rs.cell(k, 1, lib).font = F_NORMAL
        c = rs.cell(k, 2, val)
        c.font, c.number_format = F_GRAS, fmt
    k = r + len(chiffres) + 1
    if len(a_verifier) > len(affichees):
        rs.cell(k, 1, f'Seules les {len(affichees):,} paires les plus probables sont listées '
                      f'(option --max-verif pour en afficher plus).'.replace(',', ' ')).font = F_GRIS
        k += 1
    if dates_defaut:
        rs.cell(k, 1, 'Dates « par défaut » détectées (comptées comme année seulement) : '
                + ', '.join(f'{j} ({pc:.1%})'.replace('.', ',') for j, pc in dates_defaut[:4])).font = F_GRIS
        k += 1
    nb_ign = len(ignores)
    if nb_ign:
        rs.cell(k, 1, f'{nb_ign} ligne(s) ignorée(s) : nom vide.').font = F_GRIS

    r = 13
    rs.cell(r, 1, 'Vos vérifications — se remplissent quand vous répondez dans les feuilles').font = F_GRAS
    _ligne_tableau(rs, r + 1, ['Feuille', 'À traiter', 'Déjà vérifiés', 'Oui', 'Non', 'En partie', 'Précision'])
    plages = [('À fusionner', len(groupes), f"'À fusionner'!$C$2:$C${fin_f}"),
              ('À vérifier', len(affichees), f"'À vérifier'!$C$2:$C${fin_v}")]
    for k, (nom, total, rng) in enumerate(plages, r + 2):
        rs.cell(k, 1, nom).font = F_NORMAL
        rs.cell(k, 2, total)
        rs.cell(k, 3, f'=COUNTA({rng})')
        rs.cell(k, 4, f'=COUNTIF({rng},"Oui")')
        rs.cell(k, 5, f'=COUNTIF({rng},"Non")')
        rs.cell(k, 6, f'=COUNTIF({rng},"En partie")')
        rs.cell(k, 7, f'=IF(D{k}+E{k}=0,"—",D{k}/(D{k}+E{k}))').number_format = '0%'
    rs.cell(r + 4, 1, 'Précision = part des alertes qui étaient de vrais doublons. Visez ≥ 95 % sur « À fusionner » : '
                      "c'est ce niveau qui bloquerait la création à l'accueil.").font = F_GRIS

    r = 20
    rs.cell(r, 1, 'Ce que vous avez à faire').font = F_GRAS
    aide = [
        'Ouvrez « À fusionner ». Pour chaque groupe, comparez les fiches, puis répondez Oui / Non dans la case jaune.',
        'La fiche « ★ GARDER » est une suggestion (la plus complète, puis la plus ancienne) : changez-la si besoin.',
        'Ouvrez « À vérifier » : pour chaque paire, répondez Oui (même personne) ou Non.',
        'Revenez ici : la précision se calcule toute seule.',
        'Fusion dans l\'application : les « Oui » de « À fusionner » sont les fiches à fusionner dans la fiche ★.',
    ]
    for k, t in enumerate(aide, r + 1):
        rs.cell(k, 1, f'{k - r}. {t}').font = F_NORMAL

    r = 27
    rs.cell(r, 1, 'Comment lire le score').font = F_GRAS
    lire = [
        f'Score ≥ {D.SEUIL_BLOCAGE} % : quasi certain — l\'application bloquerait la création de la fiche.',
        f'Score {D.SEUIL_AVERTISSEMENT}–{D.SEUIL_BLOCAGE - 1} % : à vérifier — l\'application demanderait une confirmation.',
        '≈ 1985 dans une colonne Naissance = seule l\'année est connue (la date exacte n\'est pas fiable).',
        'Colonnes utilisées : ' + ', '.join(f'{k} ← « {entetes[v]} »' for k, v in cols.items()),
    ]
    for k, t in enumerate(lire, r + 1):
        rs.cell(k, 1, '• ' + t).font = F_NORMAL

    if simulation:
        _feuille_simulation(wb, simulation)
        rs.cell(r + len(lire) + 2, 1, 'Test de rappel automatique : voir la feuille « Test d\'efficacité ».').font = F_GRAS

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title != 'Résumé'
    wb.save(chemin)


def _ligne_tableau(ws, ligne, titres):
    for c, t in enumerate(titres, 1):
        cell = ws.cell(ligne, c, t)
        cell.font, cell.fill = F_ENTETE, R_ENTETE
        cell.alignment = Alignment(horizontal='center', wrap_text=True)


def _feuille_simulation(wb, simulation):
    resultats, exemples = simulation
    ws = wb.create_sheet('Test d\'efficacité', 1)
    ws.column_dimensions['A'].width = 44
    for col in 'BCDEFG':
        ws.column_dimensions[col].width = 14
    ws['A1'] = 'Test de rappel : de faux doublons fabriqués à partir de vos vraies fiches'
    ws['A1'].font = F_TITRE
    ws['A2'] = ('Chaque ligne : on copie une vraie fiche, on la déforme comme le ferait une erreur de saisie, '
                'et on regarde si l\'algorithme la rattache à l\'original.')
    ws['A2'].font = F_GRIS
    _ligne_tableau(ws, 4, ['Type d\'erreur simulée', 'Cas testés', 'Bloqués', 'Confirmation demandée',
                           'Manqués', 'Détectés', 'Score moyen'])
    k = 5
    for t, r in resultats.items():
        vals = [t, r['total'], r['bloque'], r['alerte'], r['manque'], f'=IFERROR((C{k}+D{k})/B{k},0)',
                sum(r['scores']) / len(r['scores']) / 100]
        for c, v in enumerate(vals, 1):
            ws.cell(k, c, v).font = F_NORMAL
        ws.cell(k, 6).number_format = '0%'
        ws.cell(k, 7).number_format = '0%'
        k += 1
    ws.cell(k, 1, 'Total').font = F_GRAS
    for c, col in enumerate('BCDE', 2):
        ws.cell(k, c, f'=SUM({col}5:{col}{k - 1})').font = F_GRAS
    ws.cell(k, 6, f'=IFERROR((C{k}+D{k})/B{k},0)').font = F_GRAS
    ws.cell(k, 6).number_format = '0%'
    ws.conditional_formatting.add(f'F5:F{k}', CellIsRule(operator='lessThan', formula=['0.8'], fill=R_ROUGE))

    if exemples:
        k += 3
        ws.cell(k, 1, 'Exemples de doublons simulés que l\'algorithme a manqués').font = F_GRAS
        _ligne_tableau(ws, k + 1, ['Type d\'erreur', 'Score', 'Fiche d\'origine', '', 'Copie déformée'])
        for n, (t, a, b, s) in enumerate(exemples, k + 2):
            ws.cell(n, 1, t).font = F_NORMAL
            ws.cell(n, 2, s / 100).number_format = '0%'
            ws.cell(n, 3, f"{a['nom']} {a['prenoms']} · {a['affichage_naissance']} · {a['telephone']}").font = F_NORMAL
            ws.cell(n, 5, f"{b['nom']} {b['prenoms']} · "
                          f"{b['date_naissance'].strftime('%d/%m/%Y') if b.get('date_naissance') else ''} · "
                          f"{b['telephone']}").font = F_NORMAL


# ── Programme principal ─────────────────────────────────────────────────────

def executer(fichier, sortie=None, feuille=None, seuil=None, simuler_n=0, forcees=None,
             date_ref=None, journal=print, progression=None, max_verif=20000):
    """Lit, analyse, écrit le rapport. Utilisé par la fenêtre et par la ligne de commande."""
    seuil = D.SEUIL_AVERTISSEMENT if seuil is None else seuil
    ref = date_ref or date.today()
    journal('Lecture du fichier…')
    patients, ignores, entetes, cols = lire_fichier(fichier, feuille, forcees or {}, ref)
    journal(f'{len(patients)} patients lus. Colonnes : '
            + ', '.join(f'{k} = « {entetes[v]} »' for k, v in cols.items()))
    dates_defaut = marquer_dates_par_defaut(patients)
    if dates_defaut:
        journal('Dates « par défaut » détectées, comptées comme année seulement : '
                + ', '.join(f'{j} ({pc:.0%})' for j, pc in dates_defaut[:4]))

    journal('Recherche des doublons…')
    paires, groupes, a_verifier = analyser(patients, seuil, progression)
    en_trop = sum(len(g['membres']) - 1 for g in groupes)
    bilan = (f'{len(groupes)} groupe(s) quasi certain(s) à fusionner : {en_trop} fiche(s) en trop, '
             f'soit {en_trop / max(len(patients), 1):.1%} de la base. '
             f'{len(a_verifier)} paire(s) douteuse(s) à vérifier.')
    journal(bilan)

    simulation = None
    if simuler_n:
        journal(f'Test d\'efficacité : {simuler_n} faux doublons…')
        simulation = simuler(patients, simuler_n)
        if simulation:
            tot = sum(r['total'] for r in simulation[0].values())
            ok = sum(r['bloque'] + r['alerte'] for r in simulation[0].values())
            bilan += f'\nTest : {ok}/{tot} faux doublons retrouvés ({ok / tot:.0%}).'
            journal(f'Test : {ok}/{tot} faux doublons retrouvés ({ok / tot:.0%}).')

    sortie = sortie or os.path.join(
        os.path.dirname(os.path.abspath(fichier)),
        'rapport_doublons_' + os.path.splitext(os.path.basename(fichier))[0] + '.xlsx')
    journal('Écriture du rapport Excel…')
    try:
        ecrire_rapport(sortie, fichier, patients, ignores, cols, entetes, paires, groupes, a_verifier,
                       simulation, dates_defaut, max_verif)
    except PermissionError:
        raise PermissionError(f'Impossible d\'écrire {sortie}.\nLe rapport est-il déjà ouvert dans Excel ? '
                              f'Fermez-le et relancez.')
    journal(f'Rapport écrit : {sortie}')
    return sortie, bilan


def ouvrir(chemin):
    import subprocess
    if sys.platform.startswith('win'):
        os.startfile(chemin)  # noqa
    elif sys.platform == 'darwin':
        subprocess.Popen(['open', chemin])
    else:
        subprocess.Popen(['xdg-open', chemin])


def console():
    """Mode texte : même analyse que la fenêtre, quand tkinter n'est pas installé."""
    print("Analyse des doublons patients — SEGHO-WALÉ")
    print("(mode texte : le module tkinter n'est pas installé ; pour avoir la "
          "fenêtre :\n     sudo apt install python3-tk)\n")

    dossier = os.getcwd()
    trouves = sorted(f for f in os.listdir(dossier)
                     if f.lower().endswith(('.xlsx', '.xlsm'))
                     and not f.startswith(('~$', 'rapport_doublons_')))
    if trouves:
        print('Fichiers Excel de ce dossier :')
        for i, f in enumerate(trouves, 1):
            print(f'  {i}. {f}')
        print()

    try:
        rep = input('Fichier Excel (numéro ci-dessus, ou chemin complet) : ').strip().strip('"\'')
    except EOFError:
        print("Aucune réponse possible (pas de terminal). Utilisez :\n"
              "    python3 tester_doublons_excel.py mon_fichier.xlsx --simuler 300")
        sys.exit(1)
    if rep.isdigit() and trouves and 1 <= int(rep) <= len(trouves):
        fichier = os.path.join(dossier, trouves[int(rep) - 1])
    else:
        fichier = os.path.expanduser(rep)
    if not os.path.isfile(fichier):
        print(f'Fichier introuvable : {fichier}')
        sys.exit(1)

    rep = input("Tester aussi l'efficacité ? nombre de faux doublons simulés "
                "[300, 0 = non] : ").strip()
    try:
        n = int(rep) if rep else 300
    except ValueError:
        n = 300
    print()

    def barre(i, total):
        print(f'\r  {i}/{total} fiches comparées', end='' if i < total else '\n', flush=True)

    try:
        sortie, bilan = executer(fichier, simuler_n=n, progression=barre)
    except Exception as e:  # noqa
        print(f'\nERREUR : {e}')
        sys.exit(1)
    if input('\nOuvrir le rapport maintenant ? [o/N] ').strip().lower().startswith('o'):
        ouvrir(sortie)


def fenetre():
    """Petite fenêtre : choisir le fichier, lancer, suivre l'avancement, ouvrir le rapport."""
    import queue
    import threading
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk
    except ImportError:
        console()   # pas d'interface graphique sur ce système : mode texte
        return

    root = tk.Tk()
    root.title('Analyse des doublons patients — SEGHO-WALÉ')
    root.geometry('680x430')
    root.minsize(560, 380)
    file_msgs = queue.Queue()
    etat = {'rapport': None}

    cadre = ttk.Frame(root, padding=16)
    cadre.pack(fill='both', expand=True)
    cadre.columnconfigure(0, weight=1)

    ttk.Label(cadre, text='Fichier Excel des patients (modèle d\'importation)',
              font=('Segoe UI', 10, 'bold')).grid(row=0, column=0, columnspan=2, sticky='w')
    chemin = tk.StringVar()
    ttk.Entry(cadre, textvariable=chemin).grid(row=1, column=0, sticky='ew', pady=(4, 10))

    def parcourir():
        f = filedialog.askopenfilename(
            title='Choisir le fichier Excel des patients',
            filetypes=[('Fichiers Excel', '*.xlsx *.xlsm'), ('Tous les fichiers', '*.*')])
        if f:
            chemin.set(f)
    ttk.Button(cadre, text='Parcourir…', command=parcourir).grid(row=1, column=1, padx=(8, 0), pady=(4, 10))

    tester = tk.BooleanVar(value=True)
    nb = tk.IntVar(value=300)
    ligne_test = ttk.Frame(cadre)
    ligne_test.grid(row=2, column=0, columnspan=2, sticky='w')
    ttk.Checkbutton(ligne_test, text='Tester aussi l\'efficacité avec', variable=tester).pack(side='left')
    ttk.Spinbox(ligne_test, from_=50, to=5000, increment=50, width=6, textvariable=nb).pack(side='left', padx=4)
    ttk.Label(ligne_test, text='faux doublons simulés').pack(side='left')

    bouton = ttk.Button(cadre, text='Lancer l\'analyse')
    bouton.grid(row=3, column=0, columnspan=2, sticky='w', pady=12)

    barre = ttk.Progressbar(cadre, mode='determinate', maximum=100)
    barre.grid(row=4, column=0, columnspan=2, sticky='ew')
    texte = tk.Text(cadre, height=9, wrap='word', relief='flat', background=root.cget('background'))
    texte.grid(row=5, column=0, columnspan=2, sticky='nsew', pady=(10, 6))
    cadre.rowconfigure(5, weight=1)
    texte.configure(state='disabled')

    bas = ttk.Frame(cadre)
    bas.grid(row=6, column=0, columnspan=2, sticky='e')
    b_rapport = ttk.Button(bas, text='Ouvrir le rapport', state='disabled',
                           command=lambda: ouvrir(etat['rapport']))
    b_dossier = ttk.Button(bas, text='Ouvrir le dossier', state='disabled',
                           command=lambda: ouvrir(os.path.dirname(etat['rapport'])))
    b_dossier.pack(side='right')
    b_rapport.pack(side='right', padx=8)

    def ecrire(msg):
        texte.configure(state='normal')
        texte.insert('end', msg + '\n')
        texte.see('end')
        texte.configure(state='disabled')

    def travail(f, n):
        try:
            sortie, bilan = executer(
                f, simuler_n=n,
                journal=lambda m: file_msgs.put(('log', m)),
                progression=lambda i, total: file_msgs.put(('prog', i * 100 / total)))
            file_msgs.put(('fin', (sortie, bilan)))
        except Exception as e:  # noqa
            file_msgs.put(('erreur', str(e)))

    def lancer():
        f = chemin.get().strip()
        if not f or not os.path.isfile(f):
            messagebox.showwarning('Fichier manquant', 'Choisissez d\'abord le fichier Excel avec « Parcourir… ».')
            return
        try:
            n = int(nb.get()) if tester.get() else 0
        except (tk.TclError, ValueError):
            n = 300
        bouton.configure(state='disabled')
        b_rapport.configure(state='disabled')
        b_dossier.configure(state='disabled')
        barre['value'] = 0
        texte.configure(state='normal')
        texte.delete('1.0', 'end')
        texte.configure(state='disabled')
        threading.Thread(target=travail, args=(f, n), daemon=True).start()

    bouton.configure(command=lancer)

    def surveiller():
        try:
            while True:
                genre, val = file_msgs.get_nowait()
                if genre == 'log':
                    ecrire(val)
                elif genre == 'prog':
                    barre['value'] = val
                elif genre == 'fin':
                    barre['value'] = 100
                    etat['rapport'] = val[0]
                    bouton.configure(state='normal')
                    b_rapport.configure(state='normal')
                    b_dossier.configure(state='normal')
                    if messagebox.askyesno('Analyse terminée', val[1] + '\n\nOuvrir le rapport maintenant ?'):
                        ouvrir(val[0])
                elif genre == 'erreur':
                    bouton.configure(state='normal')
                    ecrire('ERREUR : ' + val)
                    messagebox.showerror('Erreur', val)
        except queue.Empty:
            pass
        root.after(150, surveiller)

    surveiller()
    root.mainloop()


def main():
    if len(sys.argv) == 1:
        fenetre()   # lancé sans argument (double-clic) : la fenêtre
        return
    ap = argparse.ArgumentParser(description='Analyse des doublons dans un fichier Excel de patients.')
    ap.add_argument('fichier')
    ap.add_argument('--sortie', help='Rapport à écrire (défaut : rapport_doublons_<fichier>.xlsx)')
    ap.add_argument('--feuille', help='Nom de la feuille (défaut : la première)')
    ap.add_argument('--seuil', type=int, default=D.SEUIL_AVERTISSEMENT,
                    help=f'Score minimal affiché (défaut {D.SEUIL_AVERTISSEMENT})')
    ap.add_argument('--simuler', type=int, default=0, metavar='N',
                    help='Fabriquer N faux doublons pour mesurer le taux de détection')
    ap.add_argument('--max-verif', type=int, default=20000,
                    help='Nombre maximal de paires listées dans « À vérifier » (défaut 20000)')
    ap.add_argument('--date-export', help='Date de l\'export si la colonne est un âge (JJ/MM/AAAA)')
    for champ in SYNONYMES:
        ap.add_argument(f'--col-{champ.replace("_", "-")}', dest=f'col_{champ}',
                        help=f'Nom exact de la colonne « {champ} »')
    args = ap.parse_args()

    def barre(i, n):
        print(f'\r  {i}/{n} fiches comparées', end='' if i < n else '\n', flush=True)
    executer(args.fichier, args.sortie, args.feuille, args.seuil, args.simuler,
             {c: getattr(args, f'col_{c}') for c in SYNONYMES},
             lire_date(args.date_export) if args.date_export else None, progression=barre,
             max_verif=args.max_verif)


if __name__ == '__main__':
    main()
